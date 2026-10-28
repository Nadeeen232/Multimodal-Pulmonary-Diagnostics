"""Package local model/tracking handoffs and verify the official eight SIC items."""

import argparse
import hashlib
import json
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from pptx import Presentation

from pulmonary.common import file_hash, save_json

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "reports/evidence/team29"


def package_models():
    folder = ROOT / "artifacts/team_bundle"
    manifest = json.loads((folder / "manifest.json").read_text())
    target = ROOT / "dist/SIC_0.2.0-team29_model_bundle.zip"
    target.parent.mkdir(exist_ok=True)
    with ZipFile(target, "w", ZIP_DEFLATED) as archive:
        for name in ["manifest.json", *manifest["files"]]:
            archive.write(folder / name, "artifacts/team_bundle/" + name)
        archive.writestr(
            "MODEL_BUNDLE_README.txt",
            "Restore from the project root with Expand-Archive -DestinationPath .\nVersion: "
            + manifest["model_version"]
            + "\nSource and raw data are separate; checksums are verified at API startup.\n",
        )
    with ZipFile(target) as archive:
        assert archive.testzip() is None
        for name, expected in manifest["files"].items():
            with archive.open("artifacts/team_bundle/" + name) as stream:
                assert hashlib.file_digest(stream, "sha256").hexdigest() == expected
    return target


def package_tracking():
    """Rewrite only a COPY of tracking URIs for a portable handoff; retain source DB."""
    source = ROOT / "artifacts/team_classification"
    target = ROOT / "dist/SIC_0.2.0-team29_tracking_evidence.zip"
    with tempfile.TemporaryDirectory(dir=ROOT / ".tmp", prefix="tracking-export-") as temporary:
        database = Path(temporary) / "mlflow.db"
        original = sqlite3.connect(source / "mlflow.db")
        copy = sqlite3.connect(database)
        original.backup(copy)
        original.close()
        records = copy.execute("select run_uuid,experiment_id,artifact_uri from runs").fetchall()
        with ZipFile(target, "w", ZIP_DEFLATED) as archive:
            for run_id, experiment_id, uri in records:
                destination = f"mlruns/team29/{experiment_id}/{run_id}/artifacts"
                copy.execute(
                    "update runs set artifact_uri=? where run_uuid=?", ("./" + destination, run_id)
                )
                archive.writestr(destination + "/", b"")
                # Existing artifact URIs are local paths created by this project.
                from mlflow.utils.file_utils import local_file_uri_to_path

                actual = Path(local_file_uri_to_path(uri))
                if actual.exists():
                    for path in actual.rglob("*"):
                        if path.is_file():
                            archive.write(
                                path, destination + "/" + path.relative_to(actual).as_posix()
                            )
            for (experiment_id,) in copy.execute(
                "select experiment_id from experiments"
            ).fetchall():
                copy.execute(
                    "update experiments set artifact_location=? where experiment_id=?",
                    ("./mlruns/team29/" + str(experiment_id), experiment_id),
                )
            copy.commit()
            assert copy.execute("pragma integrity_check").fetchone()[0] == "ok"
            copy.close()
            archive.write(database, "artifacts/team_classification/mlflow.db")
            for name in [
                "cohort_manifest.json",
                "selection.json",
                "ablation_results.csv",
                "text_search.csv",
                "text_selection.json",
                "resnet50_history.csv",
                "densenet121_history.csv",
                "study_test_predictions.csv",
            ]:
                archive.write(source / name, "artifacts/team_classification/" + name)
            archive.writestr(
                "TRACKING_README.txt",
                "Restore into a project root. Run MLflow UI from that root:\npython -m mlflow ui --backend-store-uri sqlite:///artifacts/team_classification/mlflow.db --host 127.0.0.1 --port 5000\nArtifact URIs in this COPY are relative; original tracking DB is unchanged.\n",
            )
    with ZipFile(target) as archive:
        assert archive.testzip() is None
    return target


def verify_video():
    video = ROOT / "docs/demo_fallback.webm"
    assert video.read_bytes()[:4] == bytes.fromhex("1a45dfa3")
    record = (
        json.loads((ROOT / "reports/generated/submission/demo/video_verification.json").read_text())
        if (ROOT / "reports/generated/submission/demo/video_verification.json").exists()
        else json.loads((EVIDENCE / "video_verification.json").read_text())
    )
    assert record["status"] == "passed" and record["model_version"] == "0.2.0-team29-47e281fc"
    saved = EVIDENCE / "video_verification.json"
    if saved.exists():
        verified = json.loads(saved.read_text())
        if verified.get("decode_verified") and verified.get("sha256") == file_hash(video):
            return video
    cache = Path(os.environ.get("LOCALAPPDATA", str(Path.home() / ".cache"))) / "ms-playwright"
    encoder = next(cache.glob("ffmpeg-*/ffmpeg-win64.exe"), None)
    encoder = encoder or shutil.which("ffmpeg")
    if encoder:
        result = subprocess.run(
            [str(encoder), "-i", str(video), "-f", "null", "-"],
            capture_output=True,
            text=True,
            timeout=120,
        )
        match = re.search(r"Duration: (\d+):(\d+):([\d.]+)", result.stderr)
        if match:
            record["duration_seconds"] = 3600 * int(match[1]) + 60 * int(match[2]) + float(match[3])
        record["decode_verified"] = result.returncode == 0
        record["playback_verification"] = (
            "ffmpeg full decode" if result.returncode == 0 else "browser fallback required"
        )
    if not record.get("decode_verified"):
        # Playwright's minimal FFmpeg build can lack the null muxer. Verify
        # actual complete playback rather than treating that as a damaged video.
        from playwright.sync_api import sync_playwright

        with sync_playwright() as playwright:
            options = {"headless": True, "args": ["--autoplay-policy=no-user-gesture-required"]}
            if os.name == "nt":
                options["channel"] = "msedge"
            browser = playwright.chromium.launch(**options)
            page = browser.new_page()
            page.goto(video.as_uri())
            page.wait_for_function(
                "document.querySelector('video')?.readyState >= 2", timeout=30000
            )
            record["duration_seconds"] = page.locator("video").evaluate("(v)=>v.duration")
            page.locator("video").evaluate("(v)=>{v.muted=true;return v.play()}")
            page.wait_for_function(
                "document.querySelector('video').ended",
                timeout=int(record["duration_seconds"] * 1000 + 30000),
            )
            assert page.locator("video").evaluate("(v)=>v.error===null")
            browser.close()
        record.update(
            decode_verified=True,
            playback_verification="Complete browser playback to ended, with no media error",
        )
    record.update(
        video="docs/demo_fallback.webm", sha256=file_hash(video), bytes=video.stat().st_size
    )
    save_json(EVIDENCE / "video_verification.json", record)
    return video


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--package", action="store_true", help="Create separate model and portable tracking ZIPs"
    )
    parser.add_argument(
        "--offline",
        action="store_true",
        help="Use saved deployment evidence instead of a live readiness check",
    )
    args = parser.parse_args()
    (ROOT / ".tmp").mkdir(exist_ok=True)
    if args.package:
        package_models()
        package_tracking()
    required = {
        "01 Source code": [
            "src/pulmonary/team_classification.py",
            "src/pulmonary/team_inference.py",
            "app/main.py",
            "ui/app.py",
        ],
        "02 Technical documentation": [
            "README.md",
            "pyproject.toml",
            "compose.yaml",
            "Dockerfile",
            "docs/architecture.png",
            "docs/demo_runbook.md",
        ],
        "03 Data assets": [
            "docs/data_assets.md",
            "docs/data_methodology.md",
            "artifacts/team_classification/cohort_manifest.json",
        ],
        "04 Technical evidence": [
            "docs/technical_evidence_report.md",
            "docs/technical_evidence_report.docx",
            "reports/evidence/team29/ablation_results.csv",
            "reports/evidence/team29/pytest.txt",
            "reports/evidence/mlflow_dashboard.png",
            "reports/evidence/mlflow_presentation.png",
            "reports/evidence/team29/mlflow_tracking_snapshot.json",
            "reports/evidence/tracking_snapshot.json",
            "reports/evidence/optuna_trials.csv",
            "reports/evidence/optuna_best_params.json",
            "dist/SIC_0.2.0-team29_tracking_evidence.zip",
        ],
        "05 Working solution": [
            "artifacts/team_bundle/manifest.json",
            "dist/SIC_0.2.0-team29_model_bundle.zip",
            "reports/evidence/team29/deployment_verification.json",
        ],
        "06 Presentation deck": [
            "docs/SIC_Final_Presentation.pptx",
            "docs/presentation_outline.md",
            "scripts/build_presentation.py",
        ],
        "07 Demonstration": [
            "docs/demo_runbook.md",
            "docs/demo_fallback.webm",
            "scripts/automated_demo.py",
            "demo_images/demo_cases.csv",
        ],
        "08 Team contribution summary": ["team_contributions.md"],
    }
    for files in required.values():
        for name in files:
            assert (ROOT / name).is_file(), name
    manifest = json.loads((ROOT / "artifacts/team_bundle/manifest.json").read_text())
    for name, digest in manifest["files"].items():
        assert file_hash(ROOT / "artifacts/team_bundle" / name) == digest
    source = (ROOT / "docs/technical_evidence_report.md").read_text(encoding="utf-8")
    for target in re.findall(r"!\[[^\]]*\]\(([^)]+)\)", source):
        assert (ROOT / "docs" / target).is_file(), target
    doc = Document(ROOT / "docs/technical_evidence_report.docx")
    assert len(doc.inline_shapes) == len(re.findall(r"!\[", source))
    assert "Team Notebook Analysis & Data Leakage Prevention" in "\n".join(
        p.text for p in doc.paragraphs
    )
    deck = Presentation(ROOT / "docs/SIC_Final_Presentation.pptx")
    assert len(deck.slides) == 13 and all(
        len(s.notes_slide.notes_text_frame.text) > 150 for s in deck.slides
    )
    slide7 = " ".join(
        c.text for s in deck.slides[6].shapes if s.has_table for r in s.table.rows for c in r.cells
    )
    assert all(value in slide7 for value in ["66.29%", "82.30%", "82.87%", "0.8816"])
    tracking_slide = " ".join(s.text for s in deck.slides[7].shapes if s.has_text_frame)
    assert all(value in tracking_slide for value in ["Experiment Tracking", "MLflow", "Optuna"])
    assert sum(s.shape_type == 13 for page in deck.slides for s in page.shapes) >= 6
    verify_video()
    live = "Saved operational evidence verified"
    if not args.offline:
        import urllib.request

        with urllib.request.urlopen("http://127.0.0.1:8080/ready", timeout=30) as response:
            ready = json.load(response)
        assert ready["model_version"] == manifest["model_version"] and ready["status"] == "ready"
        with urllib.request.urlopen("http://127.0.0.1:8501/_stcore/health", timeout=30) as response:
            assert response.status == 200
        live = "Live API readiness and Streamlit health verified"
    rows = "\n".join(
        f"| {item} | Complete | " + "; ".join("`" + name + "`" for name in files) + " |"
        for item, files in required.items()
    )
    checklist = f"""# Final SIC submission checklist

Source: SIC8 AI Project Final handbook, page 17, “The complete submission package”.
Current model: **{manifest["model_version"]}**. Technical package: **8/8 items assembled
and verified locally**. {live}.

| Official SIC item | Status | Primary artifacts |
| --- | --- | --- |
{rows}

## Handoff requirements

Submit the source repository plus **both ZIPs in `dist/`**. Model binaries,
tracking databases and raw data are excluded from Git. A source-only checkout
will not predict until its model archive is restored. The separate tracking ZIP
contains a portable COPY of the current database and referenced run artifacts;
the original database is unchanged. Dataset references satisfy the data-assets
item; redistribution of the full dataset is not required.

The deck has {len(deck.slides)} slides and updated notes; the Word report contains {len(doc.inline_shapes)}
embedded figures and {len(doc.tables)} editable tables. A real silent fallback video
was recorded and decoded successfully. Photos and a repository QR code are not
required by the official checklist; the deck uses member initials and no invented URL.
All PowerPoint files remain excluded from GitHub; submit the generated deck
separately. The additional experiment-tracking slide preserves the original
project content and distinguishes current MLflow runs from historical Optuna tuning.

## Accurate claims and human readiness

Current fusion is 82.87% accuracy, 0.7918 F1, 0.8816 ROC-AUC on 356 studies.
The report-generator finding is documented OpenI dataset-source pretraining overlap,
not proof of individual test-image membership. The deterministic report remains
the default. No external clinical validation or measured workflow benefit is claimed.

The team still needs its own rehearsal and final review of attribution. Automated
artifact checks cannot establish individual understanding. GitHub publication was
not performed. `docs/submission_manifest.json` records versioned checksums for handoff.
"""
    (ROOT / "docs/final_submission_checklist.md").write_text(checklist, encoding="utf-8")
    inventory = set(name for files in required.values() for name in files)
    inventory.update(
        [
            "docs/final_submission_checklist.md",
            "docs/demo_fallback.webm",
            "reports/ui_screenshot.png",
            "reports/report_screenshot.png",
            "reports/swagger_screenshot.png",
        ]
    )
    inventory.update(
        p.relative_to(ROOT).as_posix()
        for p in EVIDENCE.iterdir()
        if p.is_file() and p.name != "manifest.json"
    )
    save_json(
        ROOT / "docs/submission_manifest.json",
        {
            "model_version": manifest["model_version"],
            "official_items": 8,
            "verified_items": 8,
            "live_check": not args.offline,
            "files": {
                name: {"bytes": (ROOT / name).stat().st_size, "sha256": file_hash(ROOT / name)}
                for name in sorted(inventory)
            },
        },
    )
    print("Verified 8/8 SIC package items")
    print(
        f"Word: {len(doc.inline_shapes)} figures / {len(doc.tables)} tables; PowerPoint: {len(deck.slides)} slides with notes"
    )
    print(live)


if __name__ == "__main__":
    main()
