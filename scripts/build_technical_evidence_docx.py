"""Convert the evidence Markdown to Word with embedded images and editable tables.

Requires python-docx. Run: python scripts/build_technical_evidence_docx.py
"""

import argparse
import re
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs/technical_evidence_report.md"
OUTPUT = ROOT / "docs/technical_evidence_report.docx"
TOKEN = re.compile(r"(!\[[^\]]*\]\([^)]+\)|\[[^\]]+\]\([^)]+\)|\*\*[^*]+\*\*|`[^`]+`)")
IMAGE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def add_image(paragraph, target, width=6.6, max_height=4.6):
    path = SOURCE.parent / target
    with Image.open(path) as image:
        ratio = image.height / image.width
    width = min(width, max_height / ratio)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))


def inline(paragraph, text, image_width=3.1):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("!["):
            match = IMAGE.fullmatch(part)
            add_image(paragraph, match[2], width=image_width, max_height=3.4)
        elif part.startswith("["):
            match = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
            # Word supports relative hyperlinks to evidence beside the document.
            relation = paragraph.part.relate_to(
                match[2],
                "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                is_external=True,
            )
            link = OxmlElement("w:hyperlink")
            link.set(qn("r:id"), relation)
            run = OxmlElement("w:r")
            props = OxmlElement("w:rPr")
            color = OxmlElement("w:color")
            color.set(qn("w:val"), "193EB0")
            props.append(color)
            run.append(props)
            content = OxmlElement("w:t")
            content.text = match[1]
            run.append(content)
            link.append(run)
            paragraph._p.append(link)
        elif part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def add_table(document, lines):
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if all(re.fullmatch(r"[-: ]+", cell) for cell in cells):
            continue
        rows.append(cells)
    table = document.add_table(rows=0, cols=len(rows[0]))
    table.style = "Light Shading Accent 1"
    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for cell, value in zip(cells, values, strict=True):
            inline(cell.paragraphs[0], value, image_width=3.1)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(5)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5 if len(values) >= 7 else 10)
                    if index == 0:
                        run.bold = True
        if index == 0:
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
        no_split = OxmlElement("w:cantSplit")
        table.rows[-1]._tr.get_or_add_trPr().append(no_split)
    document.add_paragraph()


def main(output=OUTPUT):
    output = Path(output)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.left_margin = section.right_margin = Inches(0.65)
    section.top_margin = section.bottom_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.05
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        document.styles[name].font.color.rgb = RGBColor.from_string("193EB0")
    header = section.header.paragraphs[0]
    header.text = "Pulmonary Review Assistant | Technical Evidence"
    header.style = document.styles["Caption"]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Technical Evidence & Working Solution  |  ")
    page = OxmlElement("w:fldSimple")
    page.set(qn("w:instr"), "PAGE")
    footer._p.append(page)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if not line.strip():
            index += 1
            continue
        if line.startswith("```"):
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_after = Pt(1)
                run = paragraph.add_run(lines[index])
                run.font.name, run.font.size = "Consolas", Pt(8)
                index += 1
            index += 1
        elif line.startswith("|"):
            block = []
            while index < len(lines) and lines[index].startswith("|"):
                block.append(lines[index])
                index += 1
            add_table(document, block)
        elif line.startswith("#"):
            match = re.match(r"(#+)\s+(.*)", line)
            level = len(match[1])
            if level == 1:
                document.add_paragraph(match[2], "Title")
            else:
                paragraph = document.add_heading(match[2], level=min(level - 1, 3))
                paragraph.paragraph_format.page_break_before = False
            index += 1
        elif IMAGE.fullmatch(line.strip()):
            match = IMAGE.fullmatch(line.strip())
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            add_image(paragraph, match[2])
            caption = document.add_paragraph(match[1], "Caption")
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            index += 1
        else:
            block = []
            while (
                index < len(lines)
                and lines[index].strip()
                and not lines[index].startswith(("#", "|", "```"))
            ):
                if IMAGE.fullmatch(lines[index].strip()):
                    break
                block.append(lines[index].strip())
                index += 1
            inline(document.add_paragraph(), " ".join(block))
    document.core_properties.title = "Technical Evidence & Working Solution"
    document.core_properties.subject = "Pulmonary Review Assistant — SIC technical evidence"
    document.core_properties.author = "Pulmonary Review Assistant"
    document.core_properties.last_modified_by = "Pulmonary Review Assistant"
    document.core_properties.created = datetime(2026, 10, 28, tzinfo=timezone.utc)
    document.core_properties.modified = datetime(2026, 10, 28, tzinfo=timezone.utc)
    document.save(output)
    expected = len(IMAGE.findall(SOURCE.read_text(encoding="utf-8")))
    with ZipFile(output) as archive:
        assert archive.testzip() is None
    reopened = Document(output)
    assert len(reopened.inline_shapes) == expected
    expected_tables = sum(
        line.startswith("|") and (i == 0 or not lines[i - 1].startswith("|"))
        for i, line in enumerate(lines)
    )
    assert len(reopened.tables) == expected_tables
    print(output)
    print(
        f"Verified: {len(reopened.inline_shapes)} embedded images, {len(reopened.tables)} editable tables; valid Word archive."
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    main(parser.parse_args().output)
