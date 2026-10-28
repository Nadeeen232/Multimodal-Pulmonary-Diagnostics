"""Build the SIC deck with current results, tracking evidence and template styling."""

import argparse
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY = "193EB0"
TEAL = "00B3E3"
GRAY = "516278"
LIGHT = "F0F5FA"
WHITE = "FFFFFF"


def box(slide, x, y, w, h, color=LIGHT):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    return shape


def text(slide, value, x, y, w, h, size=20, color=NAVY, bold=False, align=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = 0
    for index, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        if align is not None:
            p.alignment = align
        p.space_after = Pt(5)
    return shape


def image(slide, path, x, y, w, h):
    path = ROOT / path
    if not path.is_file():
        raise FileNotFoundError(path)
    with Image.open(path) as im:
        ratio = im.width / im.height
    width = min(w, h * ratio)
    height = width / ratio
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - width) / 2),
        Inches(y + (h - height) / 2),
        width=Inches(width),
        height=Inches(height),
    )


def bullets(slide, items, x=0.62, y=1.5, w=9.6, h=4.6, size=21):
    return text(slide, "\n".join("• " + item for item in items), x, y, w, h, size=size)


def title(slide, value, index, width, height):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    slide._element.set("showMasterSp", "0")
    # Match the reference's 1.31-inch SIC-blue banner and white Arial heading.
    banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(width), Inches(1.31))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    banner.line.fill.background()
    text(slide, value, 0.49, 0.30, width - 0.98, 0.8, 29, WHITE, bold=True)
    rule = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(6.79), Inches(width - 1.24), Inches(0.01)
    )
    rule.fill.solid()
    rule.fill.fore_color.rgb = RGBColor.from_string("BFBFBF")
    rule.line.fill.background()
    text(
        slide,
        "Samsung Innovation Campus",
        0.63,
        6.92,
        3.5,
        0.22,
        10,
        GRAY,
    )
    text(slide, "0.2.0-team29", width - 3.2, 6.92, 2, 0.23, 10, GRAY)
    text(slide, f"{index:02d}", width - 0.75, 6.92, 0.3, 0.23, 11, NAVY, bold=True)


def table(slide, headers, rows, x, y, width, row_height=0.46, highlight_last=False):
    shape = slide.shapes.add_table(
        len(rows) + 1,
        len(headers),
        Inches(x),
        Inches(y),
        Inches(width),
        Inches(row_height * (len(rows) + 1)),
    )
    obj = shape.table
    obj.columns[0].width = Inches(width * 0.45)
    for col in list(obj.columns)[1:]:
        col.width = Inches(width * 0.55 / (len(headers) - 1))
    for r, values in enumerate([headers] + rows):
        for c, value in enumerate(values):
            cell = obj.cell(r, c)
            cell.text = str(value)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.margin_left = Inches(0.12)
            cell.margin_top = Inches(0.06)
            cell.margin_bottom = Inches(0.04)
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(
                NAVY
                if r == 0
                else "DDF3EF"
                if highlight_last and r == len(rows)
                else LIGHT
                if r % 2
                else WHITE
            )
            for p in cell.text_frame.paragraphs:
                p.font.name = "Arial"
                p.font.size = Pt(16)
                p.font.bold = r == 0 or highlight_last and r == len(rows)
                p.font.color.rgb = RGBColor.from_string(WHITE if r == 0 else NAVY)
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    return shape


def content():
    return [
        (
            "Multimodal Clinical Decision Support for Pulmonary Pathology",
            [
                "SIC – AI Track",
                "Chest X-ray + clinical context",
                "Interpretable classification and structured reporting",
            ],
            "Good morning. We are presenting Multimodal Clinical Decision Support for Pulmonary Pathology, our Samsung Innovation Campus AI project. We combine an X-ray with its clinical indication and optional comparison context. The system produces an abnormality score, a weak-label classification, an illustrative review priority, image and text explanations, and a written evidence summary. The current model is version 0.2.0-team29-47e281fc. Its held-out study accuracy is 82.87 percent. We will explain the data, architecture, evidence and deployment, including where the system remains limited. Clinician verification is required.",
        ),
        (
            "Agenda",
            [
                "Team",
                "Motivation & problem",
                "Approach & results",
                "Explainability · deployment · future",
            ],
            "We will first introduce our responsibilities and the review problem. Next, we will explain the Indiana dataset, the curated target and the multimodal architecture. We will then show the held-out comparison and how image and text attribution support the written report. Finally, we will demonstrate the Docker application and discuss the remaining research work. All primary metrics in this deck belong to the same current cohort and model version.",
        ),
        (
            "Team members",
            [],
            "Martina and Yousef own the data preparation and exploratory analysis workstream. Zeyad and Nadeen provided the model notebooks and the deep learning and fusion approach. Mohamed Ehab integrated the notebooks into Python modules, adapted training to the six-gigabyte GPU, ran the local evaluation, and connected the model to FastAPI, Docker, Streamlit, explanations and reporting. He also documented the report-generator overlap concern. These responsibilities and contributions are recorded in team_contributions.md. AI-assisted implementation and analysis were used; each member should be prepared to explain their own component and its interface to the system.",
        ),
        (
            "Motivation & problem",
            [
                "Image review requires expert attention and clinical context.",
                "Queues and fragmented information can delay review.",
                "We test whether image + indication supports transparent review.",
                "Workflow benefit is a goal; it has not been clinically measured.",
            ],
            "A reviewer needs both the radiograph and the reason for the examination. When information is fragmented or many examinations require attention, the workflow can be delayed. Our motivation is to put model evidence and clinical context together in an inspectable interface. The project evaluates predictive performance and demonstrates a working software path. It has not measured reduced turnaround time, improved patient outcomes or safe autonomous diagnosis. Review priority is illustrative rather than a clinically validated urgency score.",
        ),
        (
            "Dataset & preparation",
            [
                "Indiana IU mirror: 3,851 reports / 7,466 image records.",
                "Selected cohort: 2,371 studies / 4,543 images.",
                "UID-grouped splits; exact-image hash audit.",
                "384 × 384 padding; train-only text fitting.",
            ],
            "We acquired the Indiana University collection through the raddar Kaggle mirror. The selected experiment requires a frontal image and retains at most one frontal and one lateral view. Weak labels use core Problems and MeSH tags; uncertain and poor-quality cases are excluded. A documented strict rule moves chronic-only cases with a no-acute impression to uncertain. The resulting split has 1,660 training, 355 validation and 356 test studies, containing 3,179, 683 and 681 images. UID is a study identifier, not a proven patient identifier. The selected hash audit found no exact duplicate-image rows. Findings, impression and tags never enter classifier predictors. Impression participates only in the cohort-label rule.",
        ),
        (
            "Our approach — integrated team models",
            [
                "ResNet50 + DenseNet121 → mean image score.",
                "Indication + comparison → word/character TF-IDF + Logistic Regression.",
                "Late fusion: 70% image + 30% text.",
                "Validation selects fusion weight and the 0.49 threshold.",
            ],
            "We rebuilt the team’s selected approach as Python files because their trained classification exports were unavailable. Both CNNs start from ImageNet weights and use grayscale Lanczos letterboxing at 384 pixels followed by ImageNet normalization. Their scores are averaged. The text branch combines word and character TF-IDF with train-scaled metadata from indication and comparison. Validation selected a 70 percent image contribution and 30 percent text contribution. The balanced-accuracy threshold is 0.49. Training used batch size eight, AMP, staged fine-tuning and early stopping. Backend and disk interruptions required documented checkpoint restarts; this is a local reproduction, not a bit-exact rerun of the notebook.",
        ),
        (
            "Model comparison — held-out studies",
            [
                "356 studies / 681 images; thresholds selected on validation.",
                "Fusion: precision 83.45% · recall 75.32%.",
                "ResNet50 has higher F1; small ROC-AUC gains are not proven superiority.",
            ],
            "These are locally measured results on 356 held-out studies. For each study we average available image scores, and F1 refers to the abnormal weak-label class. Text-only accuracy is 66.29 percent, ResNet50 is 82.30, DenseNet121 is 79.21, the image ensemble is 81.74, and multimodal fusion is 82.87. Fusion ROC-AUC is 0.8816 and F1 is 0.7918. Precision is 83.45 percent and recall is 75.32 percent, with 179 true negatives, 23 false positives, 38 false negatives and 116 true positives. ResNet50 has the higher F1. We do not claim statistically reliable superiority from these small differences. The earlier v1 cohort and 10-to-1 cost policy are different and cannot be directly compared. Scores are uncalibrated.",
        ),
        (
            "Explainability & written reporting",
            [
                "Grad-CAM explains the ResNet50 branch of the final score.",
                "Text uses token occlusion, not SHAP.",
                "Rule-based report: class · impression · evidence · context · recommendation.",
            ],
            "The two image examples show actual positive and false-positive responses from the deployed model. A convincing attention map is not proof of a correct classification. Grad-CAM targets the ResNet50 branch of the fused score. Text occlusion removes one raw token at a time and measures the final-score change with image scores fixed. These are sensitivity measures, not anatomy or causal proof. The rule-based report faithfully describes this evidence and retains clinician verification. We reviewed the team’s genuine FLAN-T5 report generator and reproduced its saved reload example. However, it misses many findings in its own checks. Its XRV tag model also includes OpenI in pretraining, creating unresolved dataset-source overlap with Indiana. We have not proved that specific test images were seen. For these reasons learned generation remains experimental rather than the default report.",
        ),
        (
            "Deployment — FastAPI + Streamlit + Docker",
            [
                "Versioned, checksummed bundle loaded once.",
                "FastAPI :8080 · Streamlit :8501.",
                "Read-only model mount; CPU inference inside containers.",
                "13 focused tests + real API/UI verification passed.",
            ],
            "FastAPI loads the immutable 0.2.0 bundle once and verifies all required hashes and the preprocessing identity. The prediction and explanation endpoints share inference code and locking. The UI communicates over HTTP and does not load another model. Compose publishes the API at localhost port 8080 and the interface at 8501. Both services are healthy. Thirteen focused tests passed, and real container requests and browser uploads were verified. Missing images return validation errors; unavailable models produce readiness and inference failures rather than invented scores. The model binaries are distributed separately from Git source and must be restored before running a fresh checkout.",
        ),
        (
            "Live demonstration & recorded fallback",
            [
                "Upload image and paste its matching indication/comparison.",
                "Inspect class, score, priority, Grad-CAM and text contributions.",
                "Read/download the report; show the false-positive case.",
                "A tested silent recording is available beside this deck.",
            ],
            "We use the current demo_cases.csv and the team29 image folder. UID 3212 demonstrates a Normal prediction, UID 2348 an Abnormal prediction, and UID 2829 a false positive against the weak label. Enter both clinical fields exactly as saved, because comparison is a trained input. Their single-image scores are approximately 21.34, 90.79 and 71.06 percent. These cases were deliberately selected after evaluation to demonstrate behavior and do not estimate general performance. The main metrics aggregate views, while this interface handles one image. If the live stack is unavailable, open demo_fallback.webm beside the presentation. That real recording uses the preserved UID 1330 automation and the same deployed version.",
        ),
        (
            "Future enhancements",
            [
                "Clinician-reviewed labels and broader datasets.",
                "Verified patient grouping and external validation.",
                "Calibration and stronger text/fusion evaluation.",
                "Grounded reports with overlap audits and factuality checks.",
            ],
            "Our next steps focus on evidence quality. Better labels and verified patient grouping would define the target more reliably. External validation is needed before generalizing beyond this collection, and calibration is needed before interpreting scores as probabilities for a defined outcome. Richer text encoders or fusion methods should be evaluated against strong image-only baselines on a new evaluation design. For generation, we would need an overlap-audited image encoder, stronger grounding, factuality checks and clinician assessment. Fluent language alone is insufficient, and prospective workflow benefit remains unmeasured.",
        ),
        (
            "Thank you — questions?",
            [
                "Image + clinical context → interpretable classification → structured report.",
                "82.87% held-out accuracy · ROC-AUC 0.8816.",
                "Reproducible source files and verified Docker demonstration.",
                "Clinician verification is required.",
            ],
            "Thank you. We delivered a file-based workflow connecting audited data, comparative modeling, multimodal inference, explanations and structured reporting to a working Docker application. We kept the result and its tradeoffs visible: fusion has slightly higher accuracy and ROC-AUC, but the image-only ResNet50 has higher F1, and none of these results establish clinical readiness. Our submission includes technical documentation, evidence, versioned local model artifacts, a live demo and a recorded fallback. We welcome questions about the labels, split design, fusion, report-generation decision or deployment.",
        ),
    ]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--template",
        type=Path,
        help="Optional source template; local reference is used when present",
    )
    parser.add_argument("--output", type=Path, default=ROOT / "docs/SIC_Final_Presentation.pptx")
    args = parser.parse_args()
    records = content()
    trials = pd.read_csv(ROOT / "reports/evidence/optuna_trials.csv")
    best = trials.loc[trials.value.idxmax()]
    records.insert(
        7,
        (
            "Experiment Tracking — MLflow & Optuna",
            [
                "MLflow: parameters, losses, model evidence and frozen test metrics.",
                "Current experiment: 10 runs; 3 model families; interrupted runs retained.",
                f"Historical Optuna: 20 trials; best CV ROC-AUC {best.value:.4f}.",
                "Current tuning: 12 text settings + 21 fusion weights; validation only.",
            ],
            "We used MLflow to record training settings, epoch loss histories, checkpoint and "
            "selection evidence, and final evaluation metrics. The current team29 experiment "
            "has ten runs: seven finished, two killed and one failed. This preserves the actual "
            "interrupted training attempts. Five finished evaluations include text, image and "
            "multimodal families. Their representative accuracies are 66.29, 82.30 and 82.87 "
            "percent. The chart is generated from saved records, not a screenshot of a live UI. "
            f"The earlier v1 text pipeline used Optuna for twenty completed grouped-CV trials. "
            f"Its best trial is number {int(best.number)}, with mean CV ROC-AUC {best.value:.4f}, "
            f"C {best.params_C:.10f} and one thousand TF-IDF features. It uses a different "
            "cohort and feature pipeline, so these parameters were not used to tune the deployed "
            "fusion. The current text C is 0.3, and validation selected 70 percent image weight "
            "and a 0.49 threshold. We keep validation search separate from held-out test metrics. "
            "To inspect the actual runs, launch MLflow on port 5000 with the team_classification "
            "database, then select team29-classification-reproduction. The README contains the "
            "exact command and restoration instructions for the separate tracking archive.",
        ),
    )
    template = args.template or ROOT / "Team 1 - DL_SIC.pptx"
    logos = []
    if template.is_file():
        deck = Presentation(template)
        # Reuse only the two title-slide brand marks, not the other team's content.
        for shape in list(deck.slides[0].shapes)[:2]:
            blips = shape._element.xpath(".//a:blip")
            if blips:
                rid = blips[0].get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )
                logos.append(
                    (
                        deck.slides[0].part.related_part(rid).blob,
                        shape.left,
                        shape.top,
                        shape.width,
                        shape.height,
                    )
                )
    elif args.template is not None:
        raise FileNotFoundError(template)
    else:
        # The third-party template is excluded from Git; reproduce the same canvas.
        deck = Presentation()
        deck.slide_width, deck.slide_height = Inches(10.819444), Inches(7.5)
        for _ in records:
            deck.slides.add_slide(deck.slide_layouts[6])
    if len(deck.slides) < len(records):
        for _ in range(len(records) - len(deck.slides)):
            deck.slides.add_slide(deck.slide_layouts[6])
    for entry in list(deck.slides._sldIdLst)[len(records) :]:
        deck.part.drop_rel(entry.rId)
        deck.slides._sldIdLst.remove(entry)
    width, height = deck.slide_width / Inches(1), deck.slide_height / Inches(1)
    outline = [
        "# Current SIC presentation outline\n\nGenerated by scripts/build_presentation.py; current model 0.2.0-team29-47e281fc.\n"
    ]
    for index, (heading, items, notes) in enumerate(records, 1):
        section_index = index if index <= 7 else index - 1
        slide = deck.slides[index - 1]
        for shape in list(slide.shapes):
            slide.shapes._spTree.remove(shape._element)
        for rel in list(slide.part.rels.values()):
            if rel.reltype not in {RT.SLIDE_LAYOUT, RT.NOTES_SLIDE}:
                slide.part.drop_rel(rel.rId)
        title(slide, heading, index, width, height)
        if index == 8:
            bullets(slide, items, y=1.55, h=1.7, size=18)
            image(slide, "reports/evidence/mlflow_presentation.png", 0.6, 3.4, 9.65, 3.1)
        elif section_index == 1:
            slide.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
            for shape in list(slide.shapes):
                slide.shapes._spTree.remove(shape._element)
            for blob, left, top, logo_width, logo_height in logos:
                slide.shapes.add_picture(BytesIO(blob), left, top, logo_width, logo_height)
            text(slide, "SAMSUNG INNOVATION CAMPUS · AI TRACK", 0.65, 1.05, 8, 0.4, 13, TEAL, True)
            text(
                slide,
                "Multimodal\nClinical Decision Support\nfor Pulmonary Pathology",
                0.65,
                1.55,
                6.3,
                2.6,
                31,
                WHITE,
                True,
            )
            text(
                slide,
                "Chest X-ray + clinical context\nInterpretable evidence and structured reporting",
                0.65,
                4.35,
                6.0,
                1,
                19,
                "DCEAF5",
            )
            box(slide, 0.65, 5.75, 5.8, 0.65, TEAL)
            text(
                slide, "82.87% accuracy  ·  ROC-AUC 0.8816", 0.83, 5.91, 5.45, 0.4, 19, WHITE, True
            )
            image(slide, "demo_images/team29/02_true_positive_uid_2348.png", 7.15, 1.8, 3.1, 4.8)
            text(
                slide,
                "0.2.0-team29 · clinician verification required",
                0.65,
                7.0,
                8,
                0.25,
                11,
                "DCEAF5",
            )
        elif section_index == 3:
            roster = [
                ("M", "Martina", "Data Preparation\n& EDA"),
                ("Y", "Yousef", "Data Preparation\n& EDA"),
                ("Z", "Zeyad", "Model notebooks\n& fusion"),
                ("N", "Nadeen", "Model notebooks\n& fusion"),
                ("ME", "Mohamed Ehab", "Integration, evaluation,\nAPI, UI, XAI & reporting"),
            ]
            for number, (initial, name, role) in enumerate(roster):
                x = 0.6 + number * 2.03
                box(slide, x, 1.85, 1.86, 3.8)
                circle = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL, Inches(x + 0.47), Inches(2.15), Inches(0.92), Inches(0.92)
                )
                circle.fill.solid()
                circle.fill.fore_color.rgb = RGBColor.from_string(TEAL)
                circle.line.fill.background()
                text(slide, initial, x + 0.47, 2.39, 0.92, 0.4, 23, WHITE, True, PP_ALIGN.CENTER)
                text(slide, name, x + 0.1, 3.35, 1.66, 0.8, 18, NAVY, True, PP_ALIGN.CENTER)
                text(slide, role, x + 0.1, 4.25, 1.66, 1.1, 14, GRAY, align=PP_ALIGN.CENTER)
            text(
                slide,
                "Shared pipeline · documented responsibilities · AI-assisted development acknowledged",
                0.65,
                6.05,
                9.5,
                0.6,
                16,
                GRAY,
            )
        elif section_index == 4:
            bullets(slide, items, w=4.7, size=21)
            image(slide, "demo_images/team29/02_true_positive_uid_2348.png", 5.65, 1.65, 4.45, 4.85)
        elif section_index == 5:
            table(
                slide,
                ["Partition", "Studies", "Images"],
                [
                    ["Downloaded", 3851, 7466],
                    ["Train", 1660, 3179],
                    ["Validation", 355, 683],
                    ["Test", 356, 681],
                ],
                0.65,
                1.65,
                9.5,
                0.52,
            )
            bullets(
                slide,
                [
                    "2,371 selected studies; uncertain / poor-quality cases excluded.",
                    "Weak tag-derived target; UID does not prove patient independence.",
                    "384-pixel padding; current report text excluded from predictors.",
                ],
                y=4.65,
                h=1.7,
                size=19,
            )
        elif section_index == 6:
            bullets(slide, items, y=1.55, h=1.5, size=17)
            image(slide, "docs/architecture.png", 0.65, 3.1, 9.5, 3.55)
        elif section_index == 7:
            data = pd.read_csv(ROOT / "reports/evidence/team29/ablation_results.csv")
            labels = [
                "Text-only (LogReg)",
                "ResNet50",
                "DenseNet121",
                "Image ensemble",
                "Multimodal Fusion",
            ]
            rows = [
                [name, f"{100 * r.accuracy:.2f}%", f"{r.f1:.4f}", f"{r.roc_auc:.4f}"]
                for name, r in zip(labels, data.itertuples(), strict=True)
            ]
            table(slide, ["Model", "Accuracy", "F1", "ROC-AUC"], rows, 0.65, 1.6, 9.5, 0.53, True)
            bullets(slide, items, y=5.05, h=1.65, size=17)
        elif section_index == 8:
            image(slide, "reports/evidence/team29/xai_presentation.png", 0.6, 1.55, 9.65, 3.8)
            bullets(slide, items, y=5.5, h=1.3, size=17)
        elif section_index == 9:
            image(slide, "reports/swagger_screenshot.png", 0.6, 1.6, 4.7, 4.9)
            bullets(slide, items, x=5.65, y=1.9, w=4.5, h=4.3, size=21)
        elif section_index == 10:
            image(slide, "reports/ui_screenshot.png", 0.6, 1.6, 5.2, 5.1)
            bullets(slide, items, x=6.0, y=1.85, w=4.2, h=3.9, size=19)
            link = text(slide, "Open recorded fallback", 6.0, 6.05, 4.2, 0.5, 19, TEAL, True)
            link.text_frame.paragraphs[0].runs[0].hyperlink.address = "demo_fallback.webm"
        elif section_index == 11:
            for n, item in enumerate(items):
                x = 0.65 + (n % 2) * 4.9
                y = 1.75 + (n // 2) * 2.25
                box(slide, x, y, 4.6, 1.85)
                text(slide, f"0{n + 1}", x + 0.2, y + 0.15, 0.7, 0.5, 23, TEAL, True)
                text(slide, item, x + 0.2, y + 0.75, 4.15, 0.9, 20)
        else:
            bullets(slide, items, y=1.75, h=4.5, size=24)
        slide.notes_slide.notes_text_frame.text = notes
        screen = "\n".join("- " + item for item in items)
        if index == 3:
            screen = "| Member | Role |\n| --- | --- |\n" + "\n".join(
                f"| {name} | {role.replace(chr(10), ' ')} |" for _, name, role in roster
            )
        elif index == 7:
            screen += (
                "\n\n| Model | Accuracy | F1 | ROC-AUC |\n| --- | ---: | ---: | ---: |\n"
                + "\n".join("| " + " | ".join(row) + " |" for row in rows)
            )
        outline.extend(
            [
                f"\n## Slide {index} — {heading}\n",
                "\n### On-screen text\n",
                screen,
                f"\n\n### Speaker Notes\n\n{notes}\n",
            ]
        )
    deck.core_properties.title = "Multimodal Clinical Decision Support for Pulmonary Pathology"
    deck.core_properties.subject = "SIC AI Track - current 0.2.0-team29 submission"
    deck.core_properties.author = "Pulmonary Review Assistant"
    deck.core_properties.last_modified_by = "Pulmonary Review Assistant"
    deck.core_properties.comments = (
        "Current measured results, real evidence assets and updated speaker notes."
    )
    deck.core_properties.created = datetime(2026, 10, 28, tzinfo=timezone.utc)
    deck.core_properties.modified = datetime(2026, 10, 28, tzinfo=timezone.utc)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(args.output)
    (ROOT / "docs/presentation_outline.md").write_text("\n".join(outline), encoding="utf-8")
    check = Presentation(args.output)
    assert len(check.slides) == len(records)
    assert all(slide.notes_slide.notes_text_frame.text.strip() for slide in check.slides)
    actual = "\n".join(
        shape.text for slide in check.slides for shape in slide.shapes if shape.has_text_frame
    )
    assert "Face Detection" not in actual and "Attendance" not in actual
    assert "82.87%" in "\n".join(
        cell.text
        for shape in check.slides[6].shapes
        if shape.has_table
        for row in shape.table.rows
        for cell in row.cells
    )
    print(
        f"Created {args.output.name}: {len(records)} slides, current notes, inserted evidence assets"
    )


if __name__ == "__main__":
    main()
